import sys
import re

with open('gerador_contratos_agu.py', 'r', encoding='utf-8') as f:
    code = f.read()

# 1. ADD parse_dotacao_item
helper = '''
def parse_dotacao_item(k, v):
    import re
    if k == "Exercício":
        return (k, v, "")
    m = re.match(r'^(.*?)\s+([\d\.\-]+)$', v.strip())
    if m:
        return (k, m.group(1), m.group(2))
    return (k, v, "")

'''
if "def parse_dotacao_item" not in code:
    code = code.replace("def gerar_clausulas_agu", helper + "def gerar_clausulas_agu")

# 2. DOCX: aplicar_estilo_tabela_oficial_docx
new_docx_styler = '''def aplicar_estilo_tabela_oficial_docx(t_it, col_widths_pt=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn, nsdecls
    from docx.shared import Pt
    from docx.oxml import parse_xml
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    
    t_it.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_it.autofit = False

    # Bordas elegantes em todas as células (Azul escuro do modelo)
    tblPr = t_it._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders {} >'
        r'<w:top w:val="single" w:sz="8" w:space="0" w:color="224061"/>'
        r'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="224061"/>'
        r'<w:insideH w:val="single" w:sz="8" w:space="0" w:color="224061"/>'
        r'<w:insideV w:val="single" w:sz="8" w:space="0" w:color="224061"/>'
        r'<w:left w:val="single" w:sz="8" w:space="0" w:color="224061"/>'
        r'<w:right w:val="single" w:sz="8" w:space="0" w:color="224061"/>'
        r'</w:tblBorders>'.format(nsdecls('w'))
    )
    tblPr.append(tblBorders)

    # Margens internas
    cell_mar = parse_xml(
        r'<w:tblCellMar {} >'
        r'<w:top w:w="80" w:type="dxa"/>'
        r'<w:bottom w:w="80" w:type="dxa"/>'
        r'<w:left w:w="120" w:type="dxa"/>'
        r'<w:right w:w="120" w:type="dxa"/>'
        r'</w:tblCellMar>'.format(nsdecls('w'))
    )
    tblPr.append(cell_mar)

    if len(t_it.rows) > 0:
        t_it.rows[0]._tr.get_or_add_trPr().append(parse_xml(r'<w:tblHeader {}/>'.format(nsdecls('w'))))

    num_cols = len(t_it.columns)
    if not col_widths_pt or len(col_widths_pt) != num_cols:
        col_widths_pt = [480.0 / max(num_cols, 1)] * num_cols
    else:
        s_pts = sum(col_widths_pt)
        if s_pts > 0:
            col_widths_pt = [(w / s_pts) * 480.0 for w in col_widths_pt]

    for row in t_it.rows:
        row._tr.get_or_add_trPr().append(parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w'))))
        for ci, cell in enumerate(row.cells):
            if ci < len(col_widths_pt):
                cell.width = Pt(col_widths_pt[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.1

def aplicar_fundo_celula_docx(cell, color_hex):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)
'''
start_idx = code.find("def aplicar_estilo_tabela_oficial_docx")
end_idx = code.find("def criar_modelo_excel_itens")
if start_idx != -1 and end_idx != -1:
    code = code[:start_idx] + new_docx_styler + "\n" + code[end_idx:]

# 3. DOCX Objeto - Header shading & color
code = code.replace("RGBColor(255, 255, 255)", "RGBColor(34, 64, 97)")
code = code.replace(
    "hdr_cells[c_idx].text = str(h_text)",
    "hdr_cells[c_idx].text = str(h_text)\n                            aplicar_fundo_celula_docx(hdr_cells[c_idx], 'E2EBF4')"
)
code = code.replace(
    "hdr_cells[i_h].text = h",
    "hdr_cells[i_h].text = h\n                            aplicar_fundo_celula_docx(hdr_cells[i_h], 'E2EBF4')"
)

# 4. DOCX Dotação
old_docx_dot = '''                    t_dot = doc.add_table(rows=len(dot_items), cols=2)
                    t_dot.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for r_i, (lbl, val) in enumerate(dot_items):
                        c0 = t_dot.cell(r_i, 0)
                        c1 = t_dot.cell(r_i, 1)
                        c0.width = Inches(2.2)
                        c1.width = Inches(4.7)
                        p0 = c0.paragraphs[0]
                        r0 = p0.add_run(lbl)
                        r0.bold = True
                        r0.font.name = "Arial"
                        r0.font.size = Pt(8.5)
                        p1 = c1.paragraphs[0]
                        r1 = p1.add_run(val)
                        r1.font.name = "Arial"
                        r1.font.size = Pt(8.5)'''

new_docx_dot = '''                    t_dot = doc.add_table(rows=len(dot_items), cols=3)
                    t_dot.alignment = WD_TABLE_ALIGNMENT.CENTER
                    aplicar_estilo_tabela_oficial_docx(t_dot, col_widths_pt=[120.0, 240.0, 120.0])
                    for r_i, (lbl, val) in enumerate(dot_items):
                        k, v1, v2 = parse_dotacao_item(lbl, val)
                        c0 = t_dot.cell(r_i, 0)
                        c1 = t_dot.cell(r_i, 1)
                        c2 = t_dot.cell(r_i, 2)
                        aplicar_fundo_celula_docx(c0, 'E2EBF4')
                        
                        r0 = c0.paragraphs[0].add_run(k)
                        r0.bold = True
                        r0.font.name = "Arial"
                        r0.font.size = Pt(8.5)
                        r0.font.color.rgb = RGBColor(34, 64, 97)
                        
                        r1 = c1.paragraphs[0].add_run(v1)
                        r1.font.name = "Arial"
                        r1.font.size = Pt(8.5)
                        
                        r2 = c2.paragraphs[0].add_run(v2)
                        r2.font.name = "Arial"
                        r2.font.size = Pt(8.5)
                        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER'''
code = code.replace(old_docx_dot, new_docx_dot)

# 5. PDF Objeto
code = code.replace(
    "('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(\"#003366\"))",
    "('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(\"#E2EBF4\"))"
)
code = code.replace(
    "('GRID', (0, 0), (-1, -2), 0.5, colors.HexColor(\"#cbd5e1\"))",
    "('GRID', (0, 0), (-1, -1), 1, colors.HexColor(\"#224061\"))\n                            ,('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor(\"#224061\"))"
)
code = code.replace(
    "('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(\"#cbd5e1\"))",
    "('GRID', (0, 0), (-1, -1), 1, colors.HexColor(\"#224061\"))"
)

code = code.replace("<b><font color='white'>{h}</font></b>", "<b><font color='#224061'>{h}</font></b>")
code = code.replace("<b><font color=\"white\">{h}</font></b>", "<b><font color=\"#224061\">{h}</font></b>")
code = code.replace("<b><font color='white'>{col}</font></b>", "<b><font color='#224061'>{col}</font></b>")

# 6. PDF Dotação
old_pdf_dot_data = '''                    t_dot_data = [
                        [Paragraph(f"<b>{k}</b>", style_tbl_td), Paragraph(v, style_tbl_td)]
                        for k, v in dot_items
                    ]
                    t_dot = Table(t_dot_data, colWidths=[120, 320])
                    t_dot.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor("#f8fafc")),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                        ('TOPPADDING', (0, 0), (-1, -1), 3),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                    ]))'''

new_pdf_dot_data = '''                    t_dot_data = []
                    for k, v in dot_items:
                        pk, pv1, pv2 = parse_dotacao_item(k, v)
                        t_dot_data.append([
                            Paragraph(f"<b><font color='#224061'>{pk}</font></b>", style_tbl_td), 
                            Paragraph(pv1, style_tbl_td), 
                            Paragraph(pv2, style_tbl_td)
                        ])
                    
                    t_dot = Table(t_dot_data, colWidths=[120, 220, 100])
                    t_dot.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor("#E2EBF4")),
                        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor("#224061")),
                        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                        ('TOPPADDING', (0, 0), (-1, -1), 3),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                    ]))'''
code = code.replace(old_pdf_dot_data, new_pdf_dot_data)

with open('gerador_contratos_agu.py', 'w', encoding='utf-8') as f:
    f.write(code)

print("Modificações aplicadas!")
